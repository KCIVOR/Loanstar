from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path("outputs/loanstar-historical-record-final/LoanStar_Historical_Development_Timeline.docx")
OUT.parent.mkdir(parents=True, exist_ok=True)
NAVY, LIGHT, ALT, BORDER = "1F4E78", "D9EAF7", "F2F2F2", "D9D9D9"

def font(run, size=10.5, bold=False, color="000000"):
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)

def shade(cell, color):
    p = cell._tc.get_or_add_tcPr()
    s = OxmlElement("w:shd")
    s.set(qn("w:fill"), color)
    p.append(s)

def border(cell):
    p = cell._tc.get_or_add_tcPr()
    b = p.first_child_found_in("w:tcBorders")
    if b is None:
        b = OxmlElement("w:tcBorders")
        p.append(b)
    for side in ("top", "left", "bottom", "right"):
        x = b.find(qn(f"w:{side}"))
        if x is None:
            x = OxmlElement(f"w:{side}")
            b.append(x)
        x.set(qn("w:val"), "single")
        x.set(qn("w:sz"), "6")
        x.set(qn("w:color"), BORDER)

def set_width(cell, inches):
    cell.width = Inches(inches)
    p = cell._tc.get_or_add_tcPr()
    w = p.find(qn("w:tcW"))
    if w is None:
        w = OxmlElement("w:tcW")
        p.append(w)
    w.set(qn("w:w"), str(int(inches * 1440)))
    w.set(qn("w:type"), "dxa")

def write_cell(cell, value, *, bold=False, color="000000", align=WD_ALIGN_PARAGRAPH.LEFT, size=9):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(value)
    font(r, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    border(cell)

def heading(doc, text):
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    font(r, size=13, bold=True)

def paragraph(doc, text, lead=None):
    p = doc.add_paragraph()
    if lead:
        r = p.add_run(lead)
        font(r, bold=True)
        r = p.add_run(text[len(lead):])
        font(r)
    else:
        r = p.add_run(text)
        font(r)

def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, value in enumerate(headers):
        set_width(t.rows[0].cells[i], widths[i])
        shade(t.rows[0].cells[i], NAVY)
        write_cell(t.rows[0].cells[i], value, bold=True, color="FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER)
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            set_width(cells[i], widths[i])
            if ri % 2:
                shade(cells[i], ALT)
            write_cell(cells[i], value, bold=(i == 1), align=WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT)
    return t

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.65)
sec.bottom_margin = Inches(0.65)
sec.left_margin = Inches(0.72)
sec.right_margin = Inches(0.72)

styles = doc.styles
styles["Normal"].font.name = "Arial"
styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].paragraph_format.space_after = Pt(7)

p = doc.add_paragraph(style="Title")
p.paragraph_format.space_after = Pt(3)
r = p.add_run("LoanStar System Historical Development Timeline")
font(r, size=20, bold=True)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(14)
r = p.add_run("June to September 2026")
font(r, size=10, color="555555")
r.italic = True

heading(doc, "Project Background")
paragraph(doc, "The LoanStar System project started in the first week of June 2026. The team visited LoanStar to understand its current loan process, the work of each department, and the features needed for the new system. After the visit, the team planned the system and prepared its overall structure before development started.")

heading(doc, "Development Timeline")
paragraph(doc, "The schedule below presents the development sequence for the system modules. The loan calculator and User Acceptance Testing dates use the project dates provided.")

rows = [
    ("June 1 to 5", "Project analysis", "Visited LoanStar, studied the current loan process, and listed the system needs."),
    ("June 8 to 20", "Planning and System Architecture", "Planned the system modules, user roles, loan workflow, screens, and the information the system needed to store. Prepared the overall structure of the LoanStar System before development started."),
    ("June 22 to 27", "Borrower Portal and Leads", "Created the borrower area for registration, loan applications, document uploads, status checking, and lead tracking."),
    ("June 29 to July 4", "Intake and Loan Computation", "Created the area for walk-in applications, missing requirements, loan details, and payment information."),
    ("July 6 to 11", "Verification", "Created the feature for checking borrower information, employment details, and references."),
    ("July 13 to 18", "Committee and Negotiation", "Created the feature for reviewing loan applications, recording decisions, and handling loan discussions."),
    ("July 20 to 25", "Loan Release", "Created the feature for final loan release, release documents, and record keeping."),
    ("July 27 to August 1", "Accounting and Payments", "Created the feature for recording payments, balances, and borrower payment records."),
    ("August 3 to 6", "Collection", "Created the feature for tracking due dates, collections, and overdue accounts."),
    ("August 7 to 9", "Remedial", "Created the feature for following up delayed accounts and recording legal or collection actions."),
    ("August 10 to 12", "Reports and Dashboard", "Created reports and a management dashboard for checking loan and collection information."),
    ("August 13 to 16", "User Access and Final Integration", "Set up user access, activity records, and checked that the main modules worked together before calculator development started."),
    ("August 17 to 28", "Loan Calculator", "Developed the loan calculator and payment schedule feature."),
    ("September 22 to 25", "User Acceptance Testing", "LoanStar users tested the system using their actual work process. UAT was completed."),
]
table(doc, ["Date", "Module or Activity", "What was completed"], rows, [1.3, 1.75, 4.2])

heading(doc, "Main Modules Completed")
modules = [
    "Borrower Portal", "Leads and Intake", "Loan Computation", "Verification", "Committee and Negotiation", "Loan Release", "Accounting and Payments", "Collection", "Remedial", "Reports and Dashboard", "User Access and Activity Records", "Loan Calculator",
]
for value in modules:
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(value)
    font(r)

heading(doc, "UAT Completion")
paragraph(doc, "User Acceptance Testing was completed from September 22 to September 25, 2026. During UAT, LoanStar users checked whether the system matched their daily work. Any feedback and screenshots from this stage can be added to the final report.")

heading(doc, "Next Steps")
paragraph(doc, "After UAT, the next activities may include final corrections, deployment, user training, final documentation, and support after launch. These activities should only be marked as completed when they are finished.")

f = sec.footer.paragraphs[0]
f.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = f.add_run("LoanStar System Historical Development Timeline")
font(r, size=8, color="666666")

doc.save(OUT)
print(OUT)
