from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path("outputs/loanstar-historical-record-draft/LoanStar_Historical_Development_Record_Draft.docx")
OUT.parent.mkdir(parents=True, exist_ok=True)

NAVY, BLUE, PALE, GRAY, BORDER = "1F4E78", "4472C4", "D9EAF7", "F2F2F2", "D9D9D9"

def set_font(run, size=10.5, bold=False, color="000000"):
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)

def shade(cell, color):
    props = cell._tc.get_or_add_tcPr()
    item = OxmlElement("w:shd")
    item.set(qn("w:fill"), color)
    props.append(item)

def border(cell):
    props = cell._tc.get_or_add_tcPr()
    borders = props.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        props.append(borders)
    for name in ("top", "left", "bottom", "right"):
        node = borders.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:color"), BORDER)

def width(cell, inches):
    cell.width = Inches(inches)
    props = cell._tc.get_or_add_tcPr()
    tcw = props.find(qn("w:tcW"))
    if tcw is None:
        tcw = OxmlElement("w:tcW")
        props.append(tcw)
    tcw.set(qn("w:w"), str(int(inches * 1440)))
    tcw.set(qn("w:type"), "dxa")

def cell_text(cell, text, *, bold=False, color="000000", align=WD_ALIGN_PARAGRAPH.LEFT, size=9):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    set_font(r, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    border(cell)

def add_heading(doc, text):
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    set_font(r, size=13, bold=True)
    return p

def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        set_font(r, bold=True)
        r = p.add_run(text[len(bold_lead):])
        set_font(r)
    else:
        r = p.add_run(text)
        set_font(r)
    return p

def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, header in enumerate(headers):
        width(table.rows[0].cells[i], widths[i])
        shade(table.rows[0].cells[i], NAVY)
        cell_text(table.rows[0].cells[i], header, bold=True, color="FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER)
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            width(cells[i], widths[i])
            if row_index % 2 == 1:
                shade(cells[i], GRAY)
            cell_text(cells[i], value, bold=(i == 0), align=WD_ALIGN_PARAGRAPH.CENTER if i == 1 else WD_ALIGN_PARAGRAPH.LEFT)
    return table

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.72)
section.right_margin = Inches(0.72)

normal = doc.styles["Normal"]
normal.font.name = "Arial"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.08

title = doc.add_paragraph(style="Title")
title.paragraph_format.space_after = Pt(3)
r = title.add_run("LoanStar System Historical Development Record")
set_font(r, size=20, bold=True)

sub = doc.add_paragraph()
sub.paragraph_format.space_after = Pt(14)
r = sub.add_run("Draft for management review | June to September 2026")
set_font(r, size=10, color="555555")
r.italic = True

add_heading(doc, "Purpose")
add_body(doc, "This draft records the development history of the LoanStar System from the initial on-site analysis through User Acceptance Testing. It combines project-team information with the available repository history. Dates identified as source-control evidence come from recorded commits; the June site visit and the August 17 to 28 calculator work period are project-team records.")

add_heading(doc, "Historical Development Timeline")
timeline = [
    ("First week of June 2026", "On-site analysis", "The team visited LoanStar to understand its current loan process, workflow, and system needs.", "Project-team record"),
    ("July 6, 2026", "Repository established", "The first project code was recorded in the source repository.", "Commit eaf4f6f"),
    ("July 8, 2026", "Core module milestone", "The core LoanStar MVP was recorded, including borrower, loan processing, computation, verification, approval, release, accounting, collection, reports, user-role, and audit functions.", "Commit fc31121"),
    ("August 17 to 28, 2026", "Loan calculator development", "The team developed the loan calculator and payment-schedule work.", "Project-team record"),
    ("August 29, 2026", "Calculator evidence recorded", "Payment-frequency features and payment-schedule unification were committed to the repository.", "Commit 5fb267d"),
    ("September 22 to 25, 2026", "User Acceptance Testing", "LoanStar users tested the system. The repository records UAT data setup, UAT presentation updates, workflow fixes, and issue corrections during this period.", "Commits from Sept. 22 to 25"),
]
add_table(doc, ["Date", "Milestone", "What happened", "Evidence"], timeline, [1.25, 1.25, 3.55, 1.1])

add_heading(doc, "Module Development Record")
add_body(doc, "The July 8 core-module milestone recorded the main LoanStar modules together. The table below lists the individual modules and later improvements that are visible in the repository history.")
modules = [
    ("Borrower Portal", "Borrower registration, loan application, documents, and application tracking.", "Core module recorded July 8"),
    ("Leads and Intake", "Lead tracking and initial borrower or loan application processing.", "Core module recorded July 8"),
    ("Loan Computation", "Loan computation and payment information.", "Core module July 8; calculator work Aug. 17 to 29"),
    ("CIG Verification", "Verification of borrower details and supporting information.", "Core module recorded July 8"),
    ("Committee and Negotiation", "Loan review, approval decisions, and negotiation workflow.", "Core module July 8; later improvements Aug. 17"),
    ("Loan Release", "Release documents, release workflow, and loan handover records.", "Core module July 8; later improvements Aug. 14"),
    ("Accounting and AR", "Loan masterlist, payment posting, ledger, and receivables tracking.", "Core module July 8; UAT fixes Sept. 24"),
    ("Collection", "Collection activity, DCR, payment reconciliation, and delinquency handling.", "Core module July 8; UAT fixes Sept. 24"),
    ("Remedial", "Delinquency follow-up and demand-letter handling.", "Core module July 8; UAT fixes Sept. 24"),
    ("Reports and Dashboard", "Operational reporting and management dashboard.", "Core module July 8; reports update Aug. 17"),
    ("User Roles and Audit", "Access control, user permissions, and activity history.", "Core module recorded July 8"),
]
add_table(doc, ["Module", "Main purpose", "Development record"], modules, [1.55, 3.45, 2.15])

add_heading(doc, "UAT Record")
add_body(doc, "User Acceptance Testing took place from September 22 to September 25, 2026. During this period, the repository records UAT test-data preparation, improvements to access controls and workflow behavior, updates to UAT presentation materials, and corrections to accounting, collection, and remedial issues.")
uat = [
    ("September 22", "UAT data and workflow preparation", "Demo data was prepared for UAT testing."),
    ("September 23", "UAT support and presentation updates", "Access controls and workflow handling were improved; UAT presentation materials were updated."),
    ("September 24", "UAT issue fixes", "Accounting, collection, and remedial fixes were recorded."),
    ("September 25", "UAT completion period", "Final document-related corrections and validation updates were recorded."),
]
add_table(doc, ["Date", "Activity", "Record"], uat, [1.25, 2.25, 3.65])

add_heading(doc, "Supporting Evidence to Attach")
for item in [
    "Photo, chat message, meeting note, or presentation from the first June 2026 LoanStar visit.",
    "Screenshots of each module or dated presentation materials.",
    "UAT attendance list, test cases, client feedback, or approval message.",
    "A GitHub export or repository link showing the listed commits.",
]:
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(item)
    set_font(r)

add_heading(doc, "Next Step")
add_body(doc, "This record confirms development and UAT activity. It should not state that the full project is complete until LoanStar confirms deployment, user training, final handover, and go-live status.")

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = footer.add_run("LoanStar System Historical Development Record Draft")
set_font(r, size=8, color="666666")

doc.save(OUT)
print(OUT)
