from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path

OUT = Path("outputs/loanstar-sdlc-report/LoanStar_SDLC_Timeline_Report.docx")
OUT.parent.mkdir(parents=True, exist_ok=True)

BLUE = "1F4E78"
MID_BLUE = "4472C4"
LIGHT_BLUE = "D9EAF7"
LIGHT_GRAY = "F2F2F2"
BORDER = "D9D9D9"

def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)

def set_cell_border(cell, color=BORDER, val="single", sz="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), val)
        element.set(qn("w:sz"), sz)
        element.set(qn("w:color"), color)

def set_cell_margins(cell, top=100, start=110, bottom=100, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + m))
        if node is None:
            node = OxmlElement("w:" + m)
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")

def set_col_width(cell, inches):
    cell.width = Inches(inches)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")

def style_cell(cell, text, bold=False, color="000000", size=9.5, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    r = p.add_run(text)
    r.bold = bold
    r.font.name = "Arial"
    r._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    r._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)
    set_cell_border(cell)

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.75)
section.right_margin = Inches(0.75)

styles = doc.styles
styles["Normal"].font.name = "Arial"
styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].paragraph_format.space_after = Pt(7)
styles["Normal"].paragraph_format.line_spacing = 1.08

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
title.paragraph_format.space_after = Pt(3)
tr = title.add_run("LoanStar System SDLC Timeline Report")
tr.font.name = "Arial"
tr._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
tr._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
tr.font.size = Pt(20)
tr.font.bold = True
tr.font.color.rgb = RGBColor(0, 0, 0)

subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(15)
subtitle.add_run("Reporting period: June to September 2026").italic = True

heading = doc.add_paragraph(style="Heading 1")
heading.paragraph_format.space_before = Pt(0)
heading.paragraph_format.space_after = Pt(6)
heading.add_run("Project Status")

p = doc.add_paragraph()
p.add_run("The LoanStar System development work began in the first week of June 2026 with an on-site review and analysis of the client’s existing loan process. ")
p.add_run("User Acceptance Testing was completed from September 22 to September 25, 2026. ").bold = True
p.add_run("This report records the project timeline through completion of UAT.")

heading = doc.add_paragraph(style="Heading 1")
heading.paragraph_format.space_before = Pt(9)
heading.paragraph_format.space_after = Pt(6)
heading.add_run("SDLC Timeline")

table = doc.add_table(rows=1, cols=4)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
widths = [1.75, 1.6, 2.0, 1.9]
headers = ["SDLC Phase", "Period", "Key Activity", "Status"]
for i, header in enumerate(headers):
    cell = table.rows[0].cells[i]
    set_col_width(cell, widths[i])
    set_cell_shading(cell, BLUE)
    style_cell(cell, header, bold=True, color="FFFFFF", size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER)

rows = [
    ("Planning and Analysis", "June 1 to June 12", "On-site review of LoanStar operations and system analysis.", "Completed"),
    ("Design and Development", "June 15 to August 31", "Design and build the LoanStar loan management modules.", "Completed for UAT"),
    ("Code Review", "September 1 to September 4", "Review the implemented features and address identified issues.", "Completed"),
    ("Quality Assurance Testing", "September 7 to September 18", "Perform internal functional and regression testing before client validation.", "Completed"),
    ("User Acceptance Testing", "September 22 to September 25", "Client validation of the LoanStar System against operational requirements.", "Completed"),
]
for idx, row in enumerate(rows):
    cells = table.add_row().cells
    for i, value in enumerate(row):
        set_col_width(cells[i], widths[i])
        if idx % 2 == 1:
            set_cell_shading(cells[i], LIGHT_GRAY)
        elif row[0] == "User Acceptance Testing":
            set_cell_shading(cells[i], LIGHT_BLUE)
        style_cell(cells[i], value, bold=(i == 0 or row[0] == "User Acceptance Testing"), align=WD_ALIGN_PARAGRAPH.CENTER if i in (1, 3) else WD_ALIGN_PARAGRAPH.LEFT)

heading = doc.add_paragraph(style="Heading 1")
heading.paragraph_format.space_before = Pt(12)
heading.paragraph_format.space_after = Pt(6)
heading.add_run("Timeline Summary")

summary = doc.add_table(rows=2, cols=4)
summary.alignment = WD_TABLE_ALIGNMENT.CENTER
summary.autofit = False
summary_widths = [1.6, 1.6, 1.6, 1.6]
for i, text in enumerate(["June", "July", "August", "September"]):
    cell = summary.rows[0].cells[i]
    set_col_width(cell, summary_widths[i])
    set_cell_shading(cell, BLUE)
    style_cell(cell, text, bold=True, color="FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER)
for i, text in enumerate([
    "On-site analysis\nand planning",
    "System design\nand development",
    "Development\ncompletion",
    "Code review, QA,\nand UAT completion",
]):
    cell = summary.rows[1].cells[i]
    set_col_width(cell, summary_widths[i])
    set_cell_shading(cell, LIGHT_BLUE if i == 3 else "FFFFFF")
    style_cell(cell, text, bold=(i == 3), align=WD_ALIGN_PARAGRAPH.CENTER)

heading = doc.add_paragraph(style="Heading 1")
heading.paragraph_format.space_before = Pt(12)
heading.paragraph_format.space_after = Pt(6)
heading.add_run("Conclusion")

p = doc.add_paragraph()
p.add_run("The reporting timeline concludes with completed UAT on September 25, 2026. ").bold = True
p.add_run("Deployment, documentation, and maintenance activities may be scheduled as the next project phase if required.")

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer_run = footer.add_run("LoanStar System SDLC Timeline Report")
footer_run.font.name = "Arial"
footer_run.font.size = Pt(8)
footer_run.font.color.rgb = RGBColor(100, 100, 100)

doc.save(OUT)
print(OUT)
