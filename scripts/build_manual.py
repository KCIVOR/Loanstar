import json
import os
import sys
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

NAVY = RGBColor(0x36, 0x5F, 0x91)
ACCENT = RGBColor(0x4F, 0x81, 0xBD)
MUTED = RGBColor(0x5A, 0x5A, 0x5A)

SCREENSHOTS = {
    "Borrower portal and loan application": (
        "output/manual-screenshots/borrower-portal.png",
        "Seeded Borrower Portal workspace. The Start application button is in the upper-right area.",
    ),
    "Agent leads": (
        "output/manual-screenshots/agent-leads.png",
        "Seeded Agent Leads pipeline. The New lead button is at the upper-right.",
    ),
    "Intake and CSA application processing": (
        "output/manual-screenshots/csa-intake.png",
        "Seeded CSA Intake queue. Use the navigation and queue controls shown to open intake work.",
    ),
    "Verification and CIG review": (
        "output/manual-screenshots/cig-verification.png",
        "Seeded CIG Verification queue. Open a queued application from this module.",
    ),
    "Committee review and decisions": (
        "output/manual-screenshots/committee-review.png",
        "Seeded Committee queue. Use this module to open a file for review and voting.",
    ),
    "Release and LRA processing": (
        "output/manual-screenshots/lra-release.png",
        "Seeded LRA Release queue. Open a queued file to begin release processing.",
    ),
    "Collections accounts and borrower contact": (
        "output/manual-screenshots/collector-accounts.png",
        "Seeded Collector Assigned accounts module. Open an account from this queue to access its loan and case files.",
    ),
    "Accounting masterlist and payment review": (
        "output/manual-screenshots/ar-masterlist.png",
        "Seeded AR Masterlist. Use its search and filters to find an account.",
    ),
    "Remedial accounts": (
        "output/manual-screenshots/remedial-accounts.png",
        "Seeded Remedial recovery queue. Open an account to review remedial actions.",
    ),
    "Administration users and roles": (
        "output/manual-screenshots/admin.png",
        "Seeded Administration workspace. Use the Administration navigation to reach Users and Roles.",
    ),
}

VISUAL_WALKTHROUGHS = {
    "Borrower portal and loan application": [
        "You should see the Borrower Portal heading, your application status cards, and a Start application button at the upper right.",
        "Select Start application to open the application flow; after saving, the new or existing application appears in this same portal.",
        "Use the portal cards to return to an application, documents, or loan status when those options appear for its current stage.",
    ],
    "Agent leads": [
        "You should see the Lead pipeline heading, stage counters, a search field, and filter controls.",
        "Select New lead at the upper right to begin a lead record.",
        "Use the search and filters before opening a row so you are working on the intended lead.",
    ],
    "Intake and CSA application processing": [
        "You should see CSA intake queue, stage counters, a search field, and the New application button at the upper right.",
        "Use the Intake submenu to switch between Intake list, Leads list, and Application history.",
        "Find the intended application with search or filters, then open its row to continue intake work.",
    ],
    "Verification and CIG review": [
        "You should see CIG verification queue, the active work counters, a borrower-or-application search field, and a Filters button.",
        "Use the CIG submenu to reach denial calls, scheduled callbacks, and history when those are the task you need to complete.",
        "Open the intended queue item before recording findings or sending a file back for revision.",
    ],
    "Committee review and decisions": [
        "You should see Committee queue, counters for in-queue and decision states, a search field, and a Filters button.",
        "Use the Committee submenu to switch between Voting queue and Decision history.",
        "Search for the application, then open its row before reviewing evidence or recording a vote.",
    ],
    "Release and LRA processing": [
        "You should see LRA release queue, release-stage counters, a search field, and a Filters button.",
        "Use the Release LRA submenu to switch between Release queue and Released loans.",
        "Open a queued file before generating documents, recording signatures, or recording release.",
    ],
    "Collections accounts and borrower contact": [
        "You should see Assigned accounts, collection counters, a borrower-or-account search field, and aging or segment filters.",
        "Use the Collection submenu to reach Accounts, DCRR, DCRR history, Collector history, and Closed accounts.",
        "Find and open the correct account before entering contact activity, a payment, or a move-of-payment offer.",
    ],
    "Accounting masterlist and payment review": [
        "You should see AR masterlist, portfolio counters, a search field, and the list of accounts.",
        "Use the Accounting AR submenu to switch between Masterlist, DCRR queue, Internal transfers, History, and Rounding writeoffs.",
        "Use search and filters first, then open the intended account or queue item for the action you need.",
    ],
    "Remedial accounts": [
        "You should see Remedial recovery, counters for assigned and delinquent work, a search field, and filters.",
        "Use the Remedial submenu to move between Remedial recovery and DCRR.",
        "Open the account from the recovery list before recording a payment, sending a demand letter, or preparing DCRR work.",
    ],
    "Administration users and roles": [
        "You should see the Administration dashboard and its navigation for Users, Roles, Config, Loan Types, Checks, document templates, and audit work.",
        "Select Users or Roles in the Administration navigation rather than working from the dashboard summary.",
        "Use search and filters on the destination page before opening or changing a record.",
    ],
}

BORROWER_FORM_GUIDE = [
    ("Start with the application details", [
        "Choose the Agent when the form offers the selector, confirm Date Applied, enter the Loan Desired amount, and enter or confirm the Sales Agent.",
        "Use the loan amount and dates that apply to this application. Do not enter a payment amount in Loan Desired.",
    ]),
    ("Complete the business section when it is shown", [
        "Enter the company name, acronym, office address, landline and mobile numbers, nature of business, business email, website, TIN, fax number, branches, date established, and employee count.",
        "Use Add officer, Add stockholder, Add customer, Add supplier, Add credit reference, and Add bank account to add one row at a time. Complete the fields in the new row before adding another.",
    ]),
    ("Complete personal information", [
        "Enter legal last name and first name, middle name when applicable, civil status, present and provincial address, ownership, years of stay, place and date of birth, phone number, and dependents.",
        "When Ownership is Other, enter the required explanation in the Please specify field. The email field is read-only in the verified borrower form.",
    ]),
    ("Add employment, spouse, and income information", [
        "Enter employer or business information, contact details, current and previous employment details when applicable, and spouse information when applicable.",
        "For income, enter gross income, expenses, and net income in the matching own-income and spouse-income areas. Add Other income only when there is a real source and amount to declare.",
    ]),
    ("Add references, bank information, and documents", [
        "Use Add dependent and Add reference for each person you need to include. Enter the relationship, contact number, and address for each reference.",
        "Enter the bank name and account number in the ADB verification field. Upload each document against its named checklist item, then select Save application form.",
        "The verified SME application displayed Business Registration, representative ID, valid IDs, Mayor's Permit, TIN and community tax certificate, location sketch, bank authorization, consent form, client or supplier list, proof of transaction, bank statements, financial statements, proof of billing, and business pictures. The checklist can vary by loan type and application state.",
    ]),
]


def set_font(run, size=None, bold=None, color=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    if size: run.font.size = Pt(size)
    if bold is not None: run.font.bold = bold
    if color: run.font.color.rgb = color


def field(paragraph, code):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = code
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])


def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    run = p.add_run(text)
    set_font(run, 14 if level == 1 else 13, True, NAVY if level == 1 else ACCENT)
    p.paragraph_format.space_before = Pt(24 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(7)
    return p


def build(content, target):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1); section.bottom_margin = Inches(1)
    section.left_margin = Inches(1.25); section.right_margin = Inches(1.25)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"; normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(header.add_run(f"{content['app_name']} System User Manual"), 9, False, MUTED)

    for _ in range(7): doc.add_paragraph()
    title = doc.add_paragraph(style="Title"); title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(title.add_run(content["app_name"]), 26, True)
    subtitle = doc.add_paragraph(); subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(subtitle.add_run(content["manual_title"]), 18, False, MUTED)
    doc.add_paragraph()
    meta = doc.add_paragraph(); meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(meta.add_run(f"Version {content['version']} | {content['date']}"), 10, False, MUTED)
    doc.add_page_break()

    heading(doc, "Table of Contents")
    toc = doc.add_paragraph(); field(toc, r'TOC \\o "1-2" \\h \\z \\u')
    note = doc.add_paragraph("In Word, right-click the table of contents and choose Update Field to show page numbers.")
    for run in note.runs: set_font(run, 9, False, MUTED); run.font.italic = True
    doc.add_page_break()

    heading(doc, "Introduction")
    for para in content["introduction"].split("\n\n"):
        doc.add_paragraph(para)
    doc.add_page_break()

    for feature in content["features"]:
        heading(doc, feature["title"])
        heading(doc, "What it is for", 2)
        audience = feature.get("audience", "Users with access to this work area")
        doc.add_paragraph(f"Who uses this: {audience}. {feature['purpose']}")
        heading(doc, "Where to find it", 2); doc.add_paragraph(feature["location"])
        heading(doc, "Screenshot guide", 2)
        screenshot, default_caption = SCREENSHOTS.get(feature["title"], (None, None))
        screenshot = feature.get("screenshot_path", screenshot)
        if screenshot and os.path.isfile(screenshot):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(screenshot, width=Inches(6.0))
            caption = doc.add_paragraph(feature.get("screenshot_caption", default_caption or "Seeded role workspace"))
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in caption.runs:
                set_font(run, 9, False, MUTED); run.font.italic = True
            walkthrough = VISUAL_WALKTHROUGHS.get(feature["title"], [])
            if walkthrough:
                heading(doc, "What you should see", 2)
                for item in walkthrough:
                    doc.add_paragraph(item, style="List Number")
        else:
            guide = doc.add_table(rows=1, cols=1)
            guide.autofit = False
            cell = guide.cell(0, 0)
            cell.width = Inches(6)
            cell.text = f"Add an approved LoanStar screen capture for: {feature['location']}"
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            cell.paragraphs[0].paragraph_format.space_before = Pt(34)
            cell.paragraphs[0].paragraph_format.space_after = Pt(34)
            for run in cell.paragraphs[0].runs:
                set_font(run, 9, False, MUTED); run.font.italic = True
        if feature["title"] == "Borrower portal and loan application":
            heading(doc, "Forms and what to enter", 2)
            doc.add_paragraph(
                "Open Edit application form from the application detail. The form sections below were verified in the supplied borrower's current application. Complete only the sections that appear for your loan type."
            )
            for group, entries in BORROWER_FORM_GUIDE:
                heading(doc, group, 2)
                for entry in entries:
                    doc.add_paragraph(entry, style="List Bullet")
        heading(doc, "Steps", 2)
        for step in feature["steps"]: doc.add_paragraph(step, style="List Number")
        if feature.get("tips"):
            heading(doc, "Tips and things to watch out for", 2)
            for tip in feature["tips"]: doc.add_paragraph(tip, style="List Bullet")

    doc.add_page_break(); heading(doc, "Appendix Items Needing Clarification")
    for item in content["needs_clarification"]:
        p = doc.add_paragraph(style="List Bullet"); set_font(p.add_run(item["item"]), bold=True)
        doc.add_paragraph(item["reason"])
    footer = section.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Page "); field(footer, "PAGE"); footer.add_run(" of "); field(footer, "NUMPAGES")
    doc.save(target)


if __name__ == "__main__":
    with open(sys.argv[1], encoding="utf-8") as source:
        build(json.load(source), sys.argv[2])
